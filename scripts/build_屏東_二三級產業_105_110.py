"""屏東縣各鄉鎮市第二、三級產業場所單位數及從業員工數 105→110 年成長表。

來源：行政院主計總處 三大普查動態查詢系統（105、110 年工業及服務業普查）
  統計項：年底場所單位數、從業員工年底人數(場所單位數)；表側：鄉鎮別（總計列用縣市別）；表頭：大行業
  105 年用「大行業(第10次修訂)」、110 年用「大行業(第11次修訂)」，兩者 17 個大行業順序相同，
  僅 J 類名稱不同（105「出版、影音製作、傳播及資通訊服務業」→110「出版影音及資通訊業」）。
原則：忠實呈現，不反推、不補值。「(D)」照原樣；成長數＝110−105，任一年非數值即填「-」。
格式比照臺南農再期中報告書表4-8（每業別：從業員工數、場所單位數 × 105年、110年、成長數）。

輸出：案件/屏東_二三級產業/屏東縣二三級產業場所單位數及從業員工數_105_110.xlsx
      案件/屏東_二三級產業/原始報表/普查動態查詢_105_110.json（原始回應）
用法：python scripts/build_屏東_二三級產業_105_110.py
"""
import io
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "stat_mcp"))
import census  # noqa: E402

OUT = ROOT / "案件" / "屏東_二三級產業"
ITEMS = ["年底場所單位數", "從業員工年底人數(場所單位數)"]
KEY = {"年底場所單位數": "年底場所單位數_", "從業員工年底人數(場所單位數)": "從業員工年底人數_場所單位數_"}
SECONDARY = 5  # 前 5 個大行業（B 礦業、C 製造、D 電力燃氣、E 用水污染、F 營建）為二級，其餘為三級


def fetch(census_name, rev):
    town = census.query(census_name, ITEMS, ["鄉鎮別"], [rev])
    county = census.query(census_name, ITEMS, ["縣市別"], [rev])
    return town, county


def industries(rows):
    keys = list(rows[0].keys())
    return [k[len(KEY["年底場所單位數"]):] for k in keys if k.startswith(KEY["年底場所單位數"])]


def num(v):
    s = str(v).replace(",", "").strip()
    try:
        return int(s)
    except ValueError:
        return None


def main():
    raw = {}
    data = {}
    for yr, cname, rev in ((105, "105工業及服務", "大行業(第10次修訂)"), (110, "110工業及服務", "大行業(第11次修訂)")):
        town, county = fetch(cname, rev)
        raw[f"{yr}_鄉鎮別"], raw[f"{yr}_縣市別"] = town, county
        rows = [d for d in town["data"] if str(d["鄉鎮別"]).startswith("屏東縣")]
        tot = [d for d in county["data"] if d["縣市別"] == "屏東縣"]
        data[yr] = {"towns": rows, "total": tot[0], "inds": industries(rows)}
    (OUT / "原始報表").mkdir(parents=True, exist_ok=True)
    (OUT / "原始報表" / "普查動態查詢_105_110.json").write_text(json.dumps(raw, ensure_ascii=False, indent=1), encoding="utf-8")

    i105, i110 = data[105]["inds"], data[110]["inds"]
    assert len(i105) == len(i110) == 17, (len(i105), len(i110))
    towns = [d["鄉鎮別"].replace("屏東縣", "") for d in data[110]["towns"]]
    assert towns == [d["鄉鎮別"].replace("屏東縣", "") for d in data[105]["towns"]], "兩年鄉鎮順序不同"

    def cell(rec, item, ind):
        return rec.get(KEY[item] + ind, "")

    def raw_val(v):
        n = num(v)
        return n if n is not None else (str(v).strip() or "-")

    wb = Workbook()
    ws = wb.active
    ws.title = "105-110成長"
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(1, 1, "屏東縣各鄉鎮市第二、三級產業場所單位數及從業員工數成長表（105年、110年）").font = Font(bold=True, size=12)
    hdr = 3
    for r in range(hdr, hdr + 4):
        ws.cell(r, 1, "行政區")
    ws.merge_cells(start_row=hdr, start_column=1, end_row=hdr + 3, end_column=1)
    col = 2
    for k, (n105, n110) in enumerate(zip(i105, i110)):
        sector = "第二級產業" if k < SECONDARY else "第三級產業"
        name = n110.replace("_", "、")
        ws.cell(hdr, col, sector)
        ws.cell(hdr + 1, col, name)
        ws.merge_cells(start_row=hdr + 1, start_column=col, end_row=hdr + 1, end_column=col + 5)
        for j, lab in enumerate(["從業員工數(人)", "場所單位數(家)"]):
            ws.cell(hdr + 2, col + 3 * j, lab)
            ws.merge_cells(start_row=hdr + 2, start_column=col + 3 * j, end_row=hdr + 2, end_column=col + 3 * j + 2)
            for m, y in enumerate(["105年", "110年", "成長數"]):
                ws.cell(hdr + 3, col + 3 * j + m, y)
        col += 6
    # 第二、三級產業的大標合併
    ws.merge_cells(start_row=hdr, start_column=2, end_row=hdr, end_column=1 + 6 * SECONDARY)
    ws.merge_cells(start_row=hdr, start_column=2 + 6 * SECONDARY, end_row=hdr, end_column=col - 1)

    body = [(t, data[105]["towns"][i], data[110]["towns"][i]) for i, t in enumerate(towns)]
    body.append(("總計", data[105]["total"], data[110]["total"]))
    r = hdr + 4
    for name, a, b in body:
        ws.cell(r, 1, name)
        c = 2
        for n105, n110 in zip(i105, i110):
            for item in ("從業員工年底人數(場所單位數)", "年底場所單位數"):
                v1, v2 = raw_val(cell(a, item, n105)), raw_val(cell(b, item, n110))
                g = v2 - v1 if isinstance(v1, int) and isinstance(v2, int) else "-"
                for m, v in enumerate((v1, v2, g)):
                    x = ws.cell(r, c + m, v)
                    if isinstance(v, int):
                        x.number_format = "#,##0;-#,##0;0"
                        x.alignment = Alignment(horizontal="right")
                    else:
                        x.alignment = Alignment(horizontal="center")
                c += 3
        r += 1
    for row in ws.iter_rows(min_row=hdr, max_row=r - 1, min_col=1, max_col=col - 1):
        for x in row:
            x.border = border
            if x.row < hdr + 4:
                x.alignment = center
                x.font = Font(bold=True)
    ws.column_dimensions["A"].width = 10
    for cc in range(2, col):
        ws.column_dimensions[get_column_letter(cc)].width = 9.5
    ws.freeze_panes = ws.cell(hdr + 4, 2)
    note = r + 1
    ws.cell(note, 1, "資料來源：行政院主計總處105年及110年工業及服務業普查（三大普查動態查詢系統）；本計畫彙整。")
    ws.cell(note + 1, 1, "註：1.以場所單位為統計單位。2.「(D)」為主計總處依統計法不予揭露之數值，照原樣呈現；"
                         "兩年任一年為非數值時，成長數以「-」表示。3.105年採行業標準分類第10次修訂、110年採第11次修訂；"
                         "J類105年名稱為「出版、影音製作、傳播及資通訊服務業」。")
    out = OUT / "屏東縣二三級產業場所單位數及從業員工數_105_110.xlsx"
    wb.save(out)
    print("輸出：", out.relative_to(ROOT))
    # 給對話呈現用的精簡摘要（總計列）
    summ = []
    for n105, n110 in zip(i105, i110):
        row = {"行業": n110.replace("_", "、")}
        for item, lab in (("從業員工年底人數(場所單位數)", "員工"), ("年底場所單位數", "場所")):
            v1, v2 = raw_val(cell(data[105]["total"], item, n105)), raw_val(cell(data[110]["total"], item, n110))
            row[lab] = (v1, v2, v2 - v1 if isinstance(v1, int) and isinstance(v2, int) else "-")
        summ.append(row)
    print(json.dumps(summ, ensure_ascii=False))
    dcount = sum(1 for _, a, b in body for rec in (a, b) for k, v in rec.items() if str(v).strip() == "(D)")
    print("(D) 格數：", dcount, "鄉鎮數：", len(towns))


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    main()

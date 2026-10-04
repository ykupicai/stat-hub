"""統計資料庫（SQLite）：把下載回來的統計表（ODS/XLS/XLSX、agrstat 查詢結果 HTML）
解析成「原始格子＋長表 facts」存起來，供 MCP 工具以 SQL 查詢。

  datasets  一個來源檔或一次查詢（含網址、sha256、取得時間、本機原始檔路徑）
  sheets    每個工作表：標題、單位、表頭列數、標籤欄數、原始格子 raw_json（可逐格回查）
  facts     長表：每個數值格一列 row_label × col_label → value（含推定的 year／area；
            標籤判斷不出年度者以資料集年度填入，例如年報的「各區」列＝該年報年度）

通用解析器的假設（政府統計表的共通版型）：
  上方若干列＝標題／單位／多層表頭（合併格向右補齊、英文翻譯列略過），
  左側若干欄＝列標籤（年度或地區，合併格向下補齊），其餘為數值；
  「-」「…」「*」等符號視為資料格但不轉數值（raw 保留原字）。
解析是推定的；寫進書件前務必用 stat_db_sheet 看原始格子確認欄位對應。
"""
import hashlib
import json
import re
import sqlite3
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DB_PATH = ROOT / "資料庫" / "stat.sqlite"
RAW_DIR = ROOT / "資料庫" / "原始檔"

SCHEMA = """
CREATE TABLE IF NOT EXISTS datasets(
  dataset_id TEXT PRIMARY KEY, source TEXT, title TEXT, year INTEGER, url TEXT,
  local_path TEXT, sha256 TEXT, fetched_at TEXT, note TEXT);
CREATE TABLE IF NOT EXISTS sheets(
  sheet_id INTEGER PRIMARY KEY AUTOINCREMENT, dataset_id TEXT, sheet_name TEXT, title TEXT, unit TEXT,
  n_header INTEGER, n_label INTEGER, n_facts INTEGER, raw_json TEXT);
CREATE TABLE IF NOT EXISTS facts(
  sheet_id INTEGER, dataset_id TEXT, source TEXT, row_no INTEGER, col_no INTEGER,
  row_label TEXT, col_label TEXT, year INTEGER, area TEXT, value REAL, raw TEXT);
CREATE INDEX IF NOT EXISTS ix_facts_q ON facts(source, year, area);
CREATE INDEX IF NOT EXISTS ix_facts_sheet ON facts(sheet_id);
CREATE INDEX IF NOT EXISTS ix_sheets_ds ON sheets(dataset_id);
"""

NUM = re.compile(r"^[-－]?\d+(\.\d+)?$")
MARK = {"-", "－", "—", "…", "...", "‥", "*", "＊", "x", "X", "…(*)"}
CJK = re.compile(r"[一-鿿]")
YEAR = re.compile(r"(?:民國)?\s*(\d{2,3})\s*年")
WYEAR = re.compile(r"^(19|20)\d{2}$")
TITLE = re.compile(r"^(表|Table)\s*[\d\-\.（(]")
AREA = re.compile(r"([一-鿿]{1,6}[縣市區鄉鎮])(?:[一-鿿]{0,4}[區鄉鎮市])?$")


def connect():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB_PATH)
    con.executescript(SCHEMA)
    return con


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def _clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    if s.lower() == "nan":
        return ""
    return " ".join(s.split())


def _num(s):
    t = s.replace(",", "").replace(" ", "")
    return float(t.replace("－", "-")) if NUM.match(t) else None


def _datalike(s):
    return s != "" and (_num(s) is not None or s in MARK)


def parse_grid(rows, filled=False):
    """rows: 2D list of str。回傳 (meta, facts)。
    filled=True：合併格已依實際合併範圍填滿（read_workbook／html_table_grid 的輸出），不再猜補。
    filled=False：合併格留空，表頭以「同一上層群組內才向右補」的規則推定。"""
    rows = [[_clean(c) for c in r] for r in rows]
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]

    def is_body(r):
        rest = [c for c in r[1:] if c]
        nums = [c for c in rest if _num(c) is not None]
        dl = [c for c in rest if _datalike(c)]
        return bool(rest) and (bool(nums) and len(dl) >= 0.5 * len(rest) or len(dl) == len(rest))

    body_idx = [i for i, r in enumerate(rows) if is_body(r)]
    if not body_idx:
        return {"title": next((c for r in rows for c in r if c), ""), "unit": "", "n_header": 0, "n_label": 0}, []
    start = body_idx[0]
    title, unit, header = "", "", []
    for r in rows[:start]:
        ne = [c for c in r if c]
        if not ne:
            continue
        joined = " ".join(ne)
        if "單位" in joined or "Unit" in joined:
            unit = unit or joined
        if not title and TITLE.match(ne[0]):
            title = ne[0]
        if len(set(ne)) == 1 and not header:
            title = title or ne[0]
            continue
        header.append(list(r))
    if not filled:  # 只在上層群組相同時向右補（避免把左邊群組的子欄名補進右邊群組）
        for hi, r in enumerate(header):
            for j in range(1, width):
                if not r[j] and r[j - 1] and (hi == 0 or header[hi - 1][j] == header[hi - 1][j - 1]):
                    r[j] = r[j - 1]
    body = [rows[i] for i in body_idx]
    # 標籤欄：多數 body 列在此欄為非數值文字
    n_label = 0
    for j in range(width):
        txt = sum(1 for r in body if r[j] and not _datalike(r[j]))
        if txt > 0.5 * len(body) or (j == 0 and txt > 0):
            n_label = j + 1
        else:
            break
    n_label = max(n_label, 1)
    # 標籤欄向下補齊（垂直合併）
    prev = [""] * n_label
    labels = []
    for r in body:
        cur = []
        for j in range(n_label):
            v = r[j] or ("" if filled else (prev[j] if any(r[k] for k in range(j + 1, n_label)) or j == 0 else ""))
            cur.append(v)
        prev = cur
        labels.append(" / ".join(x for x in cur if x))
    col_labels = {}
    for j in range(n_label, width):
        parts = [h[j] for h in header if h[j] and not TITLE.match(h[j]) and "單位" not in h[j] and "Unit" not in h[j]]
        if any(CJK.search(p) for p in parts):
            parts = [p for p in parts if CJK.search(p)]
        dedup = []
        for p in parts:
            if not dedup or dedup[-1] != p:
                dedup.append(p)
        col_labels[j] = " / ".join(dedup)
    facts = []
    for (bi, r), rl in zip(zip(body_idx, body), labels):
        for j in range(n_label, width):
            c = r[j]
            if not _datalike(c):
                continue
            cl = col_labels.get(j, "")
            facts.append({"row_no": bi, "col_no": j, "row_label": rl, "col_label": cl,
                          "year": _year(rl) or _year(cl), "area": _area(rl) or _area(cl),
                          "value": _num(c), "raw": c})
    return {"title": title, "unit": unit, "n_header": len(header), "n_label": n_label}, facts


def _year(s):
    if not s:
        return None
    for part in s.split(" / "):
        m = YEAR.search(part)
        if m:
            return int(m.group(1))
        if WYEAR.match(part.strip()):
            return int(part.strip()) - 1911
    return None


def _area(s):
    if not s:
        return None
    for part in reversed(s.split(" / ")):
        p = re.sub(r"[A-Za-z0-9\.\-'’ ]+", "", part) if CJK.search(part) else part.replace(" ", "")
        if p in ("合計", "總計", "全市", "全縣"):
            continue
        m = AREA.search(p)
        if m and not YEAR.search(p):
            return m.group(0)
    return None


def save_dataset(con, dataset_id, source, title, year, url, local_path, sheets, note="", filled=True):
    """sheets: list of (sheet_name, rows)。同 dataset_id 既有資料先刪再寫（重抓即覆蓋）。"""
    cur = con.cursor()
    old = [r[0] for r in cur.execute("SELECT sheet_id FROM sheets WHERE dataset_id=?", (dataset_id,))]
    if old:
        cur.execute(f"DELETE FROM facts WHERE sheet_id IN ({','.join('?' * len(old))})", old)
        cur.execute("DELETE FROM sheets WHERE dataset_id=?", (dataset_id,))
    rel = str(Path(local_path).resolve().relative_to(ROOT)).replace("\\", "/") if local_path else None
    cur.execute("INSERT OR REPLACE INTO datasets VALUES(?,?,?,?,?,?,?,?,?)",
                (dataset_id, source, title, year, url, rel, sha256(local_path) if local_path else None,
                 datetime.now().isoformat(timespec="seconds"), note))
    total = 0
    summary = []
    for name, rows in sheets:
        meta, facts = parse_grid(rows, filled)
        cur.execute("INSERT INTO sheets(dataset_id,sheet_name,title,unit,n_header,n_label,n_facts,raw_json) VALUES(?,?,?,?,?,?,?,?)",
                    (dataset_id, name, meta["title"], meta["unit"], meta["n_header"], meta["n_label"], len(facts),
                     json.dumps(rows, ensure_ascii=False)))
        sid = cur.lastrowid
        cur.executemany("INSERT INTO facts VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                        [(sid, dataset_id, source, f["row_no"], f["col_no"], f["row_label"], f["col_label"],
                          f["year"] if f["year"] is not None else year, f["area"], f["value"], f["raw"]) for f in facts])
        total += len(facts)
        summary.append({"sheet_id": sid, "sheet": name, "title": meta["title"][:60], "facts": len(facts)})
    con.commit()
    return {"dataset_id": dataset_id, "sheets": len(sheets), "facts": total, "detail": summary}


def _fill_spans(grid, spans):
    for r, c, rs, cs in spans:
        v = grid[r][c] if r < len(grid) and c < len(grid[r]) else ""
        for dr in range(rs):
            for dc in range(cs):
                rr, cc = r + dr, c + dc
                if rr < len(grid):
                    if cc >= len(grid[rr]):
                        grid[rr] += [""] * (cc + 1 - len(grid[rr]))
                    if not grid[rr][cc]:
                        grid[rr][cc] = v
    return grid


def _trim(grid):
    grid = [list(map(str, r)) for r in grid]
    for r in grid:
        while r and not r[-1].strip():
            r.pop()
    while grid and not grid[-1]:
        grid.pop()
    return grid


def _fmt(v):
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return "" if v is None else str(v)


def _read_ods(path, max_rep=300):
    from odf import teletype
    from odf.opendocument import load
    from odf.table import Table
    doc = load(str(path))
    out = []
    for t in doc.spreadsheet.getElementsByType(Table):
        grid, spans = [], []
        r = 0
        for row in t.childNodes:
            if row.qname[1] == "table-row-group" or row.qname[1] == "table-header-rows":
                rows_iter = [x for x in row.childNodes if x.qname[1] == "table-row"]
            elif row.qname[1] == "table-row":
                rows_iter = [row]
            else:
                continue
            for rw in rows_iter:
                rrep = int(rw.getAttribute("numberrowsrepeated") or 1)
                cells, c = [], 0
                for cell in rw.childNodes:
                    if cell.qname[1] not in ("table-cell", "covered-table-cell"):
                        continue
                    crep = min(int(cell.getAttribute("numbercolumnsrepeated") or 1), max_rep)
                    txt = teletype.extractText(cell).strip() if cell.qname[1] == "table-cell" else ""
                    if cell.qname[1] == "table-cell":
                        cs = int(cell.getAttribute("numbercolumnsspanned") or 1)
                        rs = int(cell.getAttribute("numberrowsspanned") or 1)
                        if cs > 1 or rs > 1:
                            spans.append((r, c, rs, cs))
                    cells += [txt] * crep
                    c += crep
                for _ in range(min(rrep, 3) if not any(cells) else min(rrep, 50)):
                    grid.append(list(cells))
                    r += 1
        out.append((t.getAttribute("name"), _trim(_fill_spans(grid, spans))))
    return out


def _read_xls(path):
    import xlrd
    try:
        wb = xlrd.open_workbook(str(path), formatting_info=True)
    except NotImplementedError:
        wb = xlrd.open_workbook(str(path))
    out = []
    for sh in wb.sheets():
        grid = [[_fmt(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
        spans = [(rlo, clo, rhi - rlo, chi - clo) for rlo, rhi, clo, chi in getattr(sh, "merged_cells", [])]
        out.append((sh.name, _trim(_fill_spans(grid, spans))))
    return out


def _read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out = []
    for ws in wb.worksheets:
        grid = [[_fmt(v) for v in row] for row in ws.iter_rows(values_only=True)]
        spans = [(m.min_row - 1, m.min_col - 1, m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)
                 for m in ws.merged_cells.ranges]
        out.append((ws.title, _trim(_fill_spans(grid, spans))))
    return out


def read_workbook(path):
    """ODS/XLS/XLSX → [(sheet_name, rows)]，格子一律字串，合併格依實際合併範圍填滿。"""
    ext = Path(path).suffix.lower()
    return {".ods": _read_ods, ".xls": _read_xls, ".xlsx": _read_xlsx}[ext](path)


def html_table_grid(table):
    """lxml <table> → 2D list，colspan／rowspan 依實際範圍填滿。"""
    grid, pending = [], {}
    for r_i, tr in enumerate(table.xpath(".//tr")):
        row, col = [], 0
        cells = tr.xpath("./td|./th")
        ci = 0
        while ci < len(cells) or (r_i, col) in pending:
            if (r_i, col) in pending:
                row.append(pending.pop((r_i, col)))
                col += 1
                continue
            c = cells[ci]
            ci += 1
            txt = " ".join(c.text_content().split())
            cs, rs = int(c.get("colspan", 1) or 1), int(c.get("rowspan", 1) or 1)
            for k in range(cs):
                row.append(txt)
                for d in range(1, rs):
                    pending[(r_i + d, col + k)] = txt
            col += cs
        grid.append(row)
    return grid